#!/usr/bin/env python3
"""Génère `assets/waoup-modele.pptx`, le modèle PowerPoint de la charte WAOUP.

Rejouable : chaque exécution réécrit le modèle depuis zéro.

    python3 scripts/generer_modele.py            # le modèle seul
    python3 scripts/generer_modele.py --demo     # + demo/exemple-deck-charte.pptx

Sources : `assets/waoup-charte.css` (couleurs, polices, filets, verbatim) et
les règles de diapositive de `recherche/refonte2/04-powerpoint-charte.md`
(16:9, marges 24 / 18 mm, titre de douze mots, corps à 18 pt).

python-pptx ne sait pas créer une mise en page ex nihilo : on part du modèle
par défaut, on réécrit le thème, le masque et six mises en page, on supprime
les cinq autres. Les formes fixes des mises en page sont posées en XML : la
bibliothèque n'expose `add_shape` que sur les diapositives.
"""

from __future__ import annotations

import argparse
import copy
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt

# --------------------------------------------------------------------------- #
# 1. Tokens de la charte (waoup-charte.css, bloc :root)
# --------------------------------------------------------------------------- #

NOIR = "26232E"          # --noir, texte et fonds pleins
PAPIER = "FFFFFF"        # --papier, fond de page
FOND_DOUX = "F7F6F9"     # --gris-050
LIME = "E2FFA6"          # --lime, l'accent ; il ne porte jamais de texte
LIME_TRAIT = "B4DE68"    # --lime-trait, « dérivé plus foncé pour les filets »
GRIS_700 = "4A4653"
GRIS_500 = "6F6A7D"      # repères sur fond clair
GRIS_300 = "B9B5C2"      # repères et texte secondaire sur fond noir
GRIS_200 = "DCD9E1"      # --bordure
CYAN = "C7E7EF"
ARGILE = "EADFCE"
BLANC = "FFFFFF"

POLICE_TITRE = "Space Grotesk"
POLICE_TEXTE = "Inter"
POLICE_MONO = "JetBrains Mono"

# --------------------------------------------------------------------------- #
# 2. Géométrie 16:9 (13,333 × 7,5 in), marges 24 mm côtés / 18 mm haut et bas
# --------------------------------------------------------------------------- #

MM = 36000  # EMU par millimètre

LARGEUR = Emu(12192000)   # 13,333 in
HAUTEUR = Emu(6858000)    # 7,5 in
MARGE_COTE = 24 * MM
MARGE_HB = 18 * MM

CONTENU_L = MARGE_COTE
CONTENU_LARG = int(LARGEUR) - 2 * MARGE_COTE   # 10 464 000 EMU
CONTENU_BAS = int(HAUTEUR) - MARGE_HB

FILET_EP = 38100          # 3 pt, le filet lime sous le titre
FILET_LARG = 64 * MM      # 64 mm, largeur du filet de couverture de la charte

# Bandeau de titre commun aux quatre mises en page de contenu.
TITRE_H = MARGE_HB
TITRE_HAUTEUR = 1050000
FILET_Y = TITRE_H + TITRE_HAUTEUR + 55000
CORPS_Y = FILET_Y + FILET_EP + 250000
CORPS_HAUTEUR = 3800000

# Repères de bas de page, alignés sur la marge basse.
REPERE_HAUTEUR = 220000
REPERE_Y = CONTENU_BAS - REPERE_HAUTEUR

GOUTTIERE = 457200        # 48 px, l'écart de la grille à deux colonnes
COLONNE_LARG = (CONTENU_LARG - GOUTTIERE) // 2

RACINE = Path(__file__).resolve().parents[1]
MODELE = RACINE / "assets" / "waoup-modele.pptx"
DEMO = RACINE.parents[3] / "demo" / "exemple-deck-charte.pptx"


def _rgb(hexa):
    return RGBColor.from_string(hexa)


# --------------------------------------------------------------------------- #
# 3. Fabrique XML : ce que l'API python-pptx ne couvre pas sur un masque
# --------------------------------------------------------------------------- #


def _niveau_xml(rang, spec):
    """Un `<a:lvlNpPr>` : c'est lui qui décide du format du texte saisi sur la
    diapositive, pas le formatage du texte d'invite."""
    morceaux = [
        f'<a:lvl{rang}pPr marL="{spec.get("marge", 0)}" '
        f'indent="{spec.get("retrait", 0)}" algn="{spec.get("align", "l")}">',
        f'<a:lnSpc><a:spcPct val="{int(spec.get("interligne", 1.5) * 100000)}"/>'
        f"</a:lnSpc>",
        f'<a:spcBef><a:spcPts val="{int(spec.get("avant", 6) * 100)}"/></a:spcBef>',
    ]
    if spec.get("puce"):
        morceaux += [
            f'<a:buClr><a:srgbClr val="{spec.get("puce_couleur", LIME_TRAIT)}"/>'
            f"</a:buClr>",
            '<a:buSzPct val="80000"/>',
            '<a:buFont typeface="Arial"/>',
            f'<a:buChar char="{spec["puce"]}"/>',
        ]
    else:
        morceaux.append("<a:buNone/>")
    police = spec.get("police", POLICE_TEXTE)
    morceaux += [
        f'<a:defRPr sz="{int(spec.get("taille", 18) * 100)}" '
        f'b="{1 if spec.get("gras") else 0}" '
        f'i="{1 if spec.get("italique") else 0}" '
        f'spc="{spec.get("interlettrage", 0)}" '
        f'cap="{spec.get("capitales", "none")}">',
        f'<a:solidFill><a:srgbClr val="{spec.get("couleur", NOIR)}"/></a:solidFill>',
        f'<a:latin typeface="{police}"/><a:cs typeface="{police}"/>',
        "</a:defRPr>",
        f"</a:lvl{rang}pPr>",
    ]
    return "".join(morceaux)


def regler_texte(forme, niveaux, ancrage=MSO_ANCHOR.TOP, invite=None):
    """Impose le format hérité (`lstStyle`), l'ancrage et le texte d'invite."""
    tx_body = forme._element.find(qn("p:txBody"))
    ancien = tx_body.find(qn("a:lstStyle"))
    if ancien is not None:
        tx_body.remove(ancien)
    tx_body.insert(
        1,
        parse_xml(
            f'<a:lstStyle {nsdecls("a")}>'
            + "".join(_niveau_xml(i + 1, n) for i, n in enumerate(niveaux))
            + "</a:lstStyle>"
        ),
    )
    cadre = forme.text_frame
    cadre.word_wrap = True
    cadre.vertical_anchor = ancrage
    cadre.margin_left = cadre.margin_right = 0
    cadre.margin_top = cadre.margin_bottom = 0
    if invite is not None:
        cadre.text = invite  # ce que PowerPoint affiche tant que rien n'est saisi


def _identifiant(porteur):
    arbre = porteur.shapes._spTree
    return max([int(e.get("id")) for e in arbre.iter(qn("p:cNvPr"))] or [1]) + 1


def _greffer(porteur, xml):
    element = parse_xml(xml)
    porteur.shapes._spTree.append(element)
    return element


def ajouter_placeholder(porteur, nom, ph_type, idx, x, y, cx, cy):
    attr = f' type="{ph_type}"' if ph_type else ""
    _greffer(
        porteur,
        f'<p:sp {nsdecls("p", "a")}><p:nvSpPr>'
        f'<p:cNvPr id="{_identifiant(porteur)}" name="{nom}"/>'
        f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        f'<p:nvPr><p:ph{attr} idx="{idx}"/></p:nvPr></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/>'
        f'<a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
        f"<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>",
    )
    # `placeholders[i]` indexe par position sur un masque : il faut `get(idx)`.
    return porteur.placeholders.get(idx)


def _rectangle(porteur, nom, x, y, cx, cy, fond, trait=None, trait_ep=12700):
    """Un aplat, éventuellement cerné d'un filet. Jamais d'ombre, jamais de
    dégradé : `<a:effectLst/>` vide coupe tout héritage d'effet."""
    remplissage = (f'<a:solidFill><a:srgbClr val="{fond}"/></a:solidFill>'
                   if fond else "<a:noFill/>")
    contour = (
        f'<a:ln w="{trait_ep}" cap="flat" cmpd="sng" algn="ctr">'
        f'<a:solidFill><a:srgbClr val="{trait}"/></a:solidFill>'
        f'<a:prstDash val="solid"/></a:ln>'
        if trait
        else "<a:ln><a:noFill/></a:ln>"
    )
    return _greffer(
        porteur,
        f'<p:sp {nsdecls("p", "a")}><p:nvSpPr>'
        f'<p:cNvPr id="{_identifiant(porteur)}" name="{nom}"/>'
        f"<p:cNvSpPr/><p:nvPr/></p:nvSpPr>"
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/>'
        f'<a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f"{remplissage}{contour}<a:effectLst/></p:spPr>"
        f'<p:txBody><a:bodyPr lIns="0" tIns="0" rIns="0" bIns="0"/>'
        f"<a:lstStyle/><a:p/></p:txBody></p:sp>",
    )


def filet(porteur, x, y, largeur, couleur=LIME_TRAIT, epaisseur=FILET_EP,
          nom="Filet du titre"):
    """Le filet d'accent sous un titre : 3 pt, plein, sans contour."""
    return _rectangle(porteur, nom, x, y, largeur, epaisseur, couleur)


def cadre_fin(porteur, x, y, largeur, hauteur):
    """Un encadré WAOUP : cadre fin sur les quatre côtés, jamais un onglet
    coloré d'un seul côté. Il survit à la photocopie."""
    return _rectangle(porteur, "Cadre du verbatim", x, y, largeur, hauteur,
                      FOND_DOUX, trait=GRIS_200)


def _zone_de_texte(porteur, nom, x, y, cx, cy, texte, taille, couleur,
                   police, align="l", interlettrage=0, gras=True):
    return _greffer(
        porteur,
        f'<p:sp {nsdecls("p", "a")}><p:nvSpPr>'
        f'<p:cNvPr id="{_identifiant(porteur)}" name="{nom}"/>'
        f'<p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/>'
        f'<a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f"<a:noFill/><a:ln><a:noFill/></a:ln><a:effectLst/></p:spPr>"
        f'<p:txBody><a:bodyPr lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr" '
        f'wrap="square"/><a:lstStyle/>'
        f'<a:p><a:pPr algn="{align}"/><a:r>'
        f'<a:rPr lang="fr-FR" sz="{int(taille * 100)}" '
        f'b="{1 if gras else 0}" spc="{interlettrage}">'
        f'<a:solidFill><a:srgbClr val="{couleur}"/></a:solidFill>'
        f'<a:latin typeface="{police}"/><a:cs typeface="{police}"/></a:rPr>'
        f"<a:t>{texte}</a:t></a:r></a:p></p:txBody></p:sp>",
    )


def supprimer_placeholders(porteur, types_gardes):
    """Ne garde que les espaces réservés utiles. La date et le pied de page
    par défaut partent : les repères WAOUP les remplacent."""
    for forme in list(porteur.placeholders):
        if str(forme.placeholder_format.type).split()[0] not in types_gardes:
            forme._element.getparent().remove(forme._element)


def _placeholder_par_type(porteur, nom_type):
    for forme in porteur.placeholders:
        if str(forme.placeholder_format.type).startswith(nom_type):
            return forme
    raise KeyError(nom_type)


def poser_fond(porteur, couleur):
    """Aplat plein sur un masque ou une mise en page. Aucun dégradé."""
    c_sld = porteur._element.find(qn("p:cSld"))
    ancien = c_sld.find(qn("p:bg"))
    if ancien is not None:
        c_sld.remove(ancien)
    c_sld.insert(
        0,
        parse_xml(
            f'<p:bg {nsdecls("p", "a")}><p:bgPr>'
            f'<a:solidFill><a:srgbClr val="{couleur}"/></a:solidFill>'
            f"<a:effectLst/></p:bgPr></p:bg>"
        ),
    )


def poser_reperes(porteur, couleur):
    """« WAOUP » à gauche, le numéro de page à droite, JetBrains Mono 10 pt."""
    _zone_de_texte(
        porteur, "Mention WAOUP", CONTENU_L, REPERE_Y, CONTENU_LARG // 2,
        REPERE_HAUTEUR, "WAOUP", 10, couleur, POLICE_MONO, interlettrage=120,
    )
    numero = _placeholder_par_type(porteur, "SLIDE_NUMBER")
    numero.name = "Numéro de page"
    numero.left = Emu(CONTENU_L + CONTENU_LARG // 2)
    numero.top = Emu(REPERE_Y)
    numero.width = Emu(CONTENU_LARG // 2)
    numero.height = Emu(REPERE_HAUTEUR)
    regler_texte(
        numero,
        [dict(taille=10, police=POLICE_MONO, couleur=couleur, gras=True,
              interligne=1.0, avant=0, align="r")],
        ancrage=MSO_ANCHOR.MIDDLE,
    )


# --------------------------------------------------------------------------- #
# 4. Le thème : couleurs et polices de la charte dans theme1.xml
# --------------------------------------------------------------------------- #

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def ecrire_theme(presentation):
    partie = presentation.slide_masters[0].part.part_related_by(RT.THEME)
    racine = etree.fromstring(partie.blob)
    racine.set("name", "WAOUP")

    couleurs = {
        "dk1": NOIR, "lt1": PAPIER,               # texte / fond papier
        "dk2": GRIS_700, "lt2": FOND_DOUX,
        "accent1": LIME, "accent2": LIME_TRAIT,   # l'accent et son dérivé filet
        "accent3": GRIS_500, "accent4": GRIS_200,
        "accent5": CYAN, "accent6": ARGILE,
        "hlink": GRIS_700, "folHlink": GRIS_500,  # aucun bleu système
    }
    schema = racine.find(f".//{{{NS_A}}}clrScheme")
    schema.set("name", "WAOUP")
    for nom, valeur in couleurs.items():
        noeud = schema.find(f"{{{NS_A}}}{nom}")
        for enfant in list(noeud):
            noeud.remove(enfant)
        etree.SubElement(noeud, f"{{{NS_A}}}srgbClr").set("val", valeur)

    polices = racine.find(f".//{{{NS_A}}}fontScheme")
    polices.set("name", "WAOUP")
    for balise, police in (("majorFont", POLICE_TITRE), ("minorFont", POLICE_TEXTE)):
        polices.find(f"{{{NS_A}}}{balise}/{{{NS_A}}}latin").set("typeface", police)

    # Jeu de formats réduit à des aplats et des filets : aucune forme ne peut
    # hériter d'un dégradé ni d'une ombre.
    ligne = ('<a:ln w="%d" cap="flat" cmpd="sng" algn="ctr">'
             '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
             '<a:prstDash val="solid"/><a:miter lim="800000"/></a:ln>')
    aplat = '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
    format_xml = (
        f'<a:fmtScheme {nsdecls("a")} name="WAOUP">'
        f"<a:fillStyleLst>{aplat * 3}</a:fillStyleLst>"
        f"<a:lnStyleLst>{ligne % 12700}{ligne % 19050}{ligne % 28575}</a:lnStyleLst>"
        f"<a:effectStyleLst>"
        f"{'<a:effectStyle><a:effectLst/></a:effectStyle>' * 3}"
        f"</a:effectStyleLst>"
        f"<a:bgFillStyleLst>{aplat * 3}</a:bgFillStyleLst>"
        f"</a:fmtScheme>"
    )
    ancien = racine.find(f".//{{{NS_A}}}fmtScheme")
    ancien.getparent().replace(ancien, etree.fromstring(format_xml))

    partie._blob = etree.tostring(racine, xml_declaration=True,
                                  encoding="UTF-8", standalone=True)


# --------------------------------------------------------------------------- #
# 5. Le masque
# --------------------------------------------------------------------------- #

TXSTYLES = (
    f'<p:txStyles {nsdecls("p", "a")}>'
    f"<p:titleStyle>"
    f'<a:lvl1pPr algn="l"><a:lnSpc><a:spcPct val="112000"/></a:lnSpc>'
    f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:buNone/>'
    f'<a:defRPr sz="3200" b="1" spc="-64">'
    f'<a:solidFill><a:srgbClr val="{NOIR}"/></a:solidFill>'
    f'<a:latin typeface="{POLICE_TITRE}"/><a:cs typeface="{POLICE_TITRE}"/>'
    f"</a:defRPr></a:lvl1pPr>"
    f"</p:titleStyle>"
    f"<p:bodyStyle>"
    + "".join(
        f'<a:lvl{n}pPr marL="{228600 * n}" indent="-228600" algn="l">'
        f'<a:lnSpc><a:spcPct val="150000"/></a:lnSpc>'
        f'<a:spcBef><a:spcPts val="600"/></a:spcBef>'
        f'<a:buClr><a:srgbClr val="{LIME_TRAIT}"/></a:buClr>'
        f'<a:buSzPct val="80000"/><a:buFont typeface="Arial"/><a:buChar char="▪"/>'
        f'<a:defRPr sz="1800">'
        f'<a:solidFill><a:srgbClr val="{NOIR}"/></a:solidFill>'
        f'<a:latin typeface="{POLICE_TEXTE}"/><a:cs typeface="{POLICE_TEXTE}"/>'
        f"</a:defRPr></a:lvl{n}pPr>"
        for n in range(1, 6)
    )
    + "</p:bodyStyle>"
    f"<p:otherStyle>"
    f'<a:lvl1pPr><a:defRPr sz="1800">'
    f'<a:solidFill><a:srgbClr val="{NOIR}"/></a:solidFill>'
    f'<a:latin typeface="{POLICE_TEXTE}"/></a:defRPr></a:lvl1pPr>'
    f"</p:otherStyle>"
    f"</p:txStyles>"
)


def _niveaux_corps(couleur=NOIR, taille=18):
    """Corps Inter 18 pt, interligne 1,5, puces carrées lime-trait.
    Deux niveaux seulement : au-delà, la diapositive porte plus d'une idée."""
    return [
        dict(taille=taille, couleur=couleur, marge=228600, retrait=-228600,
             puce="▪", interligne=1.5, avant=10),
        dict(taille=taille - 2, couleur=couleur, marge=457200, retrait=-228600,
             puce="▪", interligne=1.5, avant=6),
    ]


def preparer_masque(masque):
    poser_fond(masque, PAPIER)
    supprimer_placeholders(masque, {"TITLE", "BODY", "SLIDE_NUMBER"})

    titre = _placeholder_par_type(masque, "TITLE")
    titre.name = "Titre"
    titre.left, titre.top = Emu(CONTENU_L), Emu(TITRE_H)
    titre.width, titre.height = Emu(CONTENU_LARG), Emu(TITRE_HAUTEUR)
    regler_texte(
        titre,
        [dict(taille=32, gras=True, police=POLICE_TITRE, interligne=1.12,
              avant=0, interlettrage=-64)],
        ancrage=MSO_ANCHOR.BOTTOM,
    )

    corps = _placeholder_par_type(masque, "BODY")
    corps.name = "Corps"
    corps.left, corps.top = Emu(CONTENU_L), Emu(CORPS_Y)
    corps.width, corps.height = Emu(CONTENU_LARG), Emu(CORPS_HAUTEUR)
    regler_texte(corps, _niveaux_corps())

    poser_reperes(masque, GRIS_500)

    anciens = masque._element.find(qn("p:txStyles"))
    if anciens is not None:
        masque._element.replace(anciens, parse_xml(TXSTYLES))
    else:
        masque._element.append(parse_xml(TXSTYLES))


# --------------------------------------------------------------------------- #
# 6. Les six mises en page
# --------------------------------------------------------------------------- #


def _bandeau_titre(mise_en_page, invite):
    """Titre ancré en bas de son cadre, filet lime dessous : la conclusion
    reste collée au filet, qu'elle tienne sur une ou deux lignes."""
    titre = _placeholder_par_type(mise_en_page, "TITLE")
    titre.name = "Titre"
    titre.left, titre.top = Emu(CONTENU_L), Emu(TITRE_H)
    titre.width, titre.height = Emu(CONTENU_LARG), Emu(TITRE_HAUTEUR)
    regler_texte(
        titre,
        [dict(taille=32, gras=True, police=POLICE_TITRE, couleur=NOIR,
              interligne=1.12, avant=0, interlettrage=-64)],
        ancrage=MSO_ANCHOR.BOTTOM,
        invite=invite,
    )
    filet(mise_en_page, CONTENU_L, FILET_Y, FILET_LARG, LIME_TRAIT)
    return titre


INVITE_TITRE = "Titre : la conclusion, douze mots au plus"


def construire_couverture(mise_en_page):
    mise_en_page.name = "Couverture"
    poser_fond(mise_en_page, NOIR)
    supprimer_placeholders(
        mise_en_page, {"CENTER_TITLE", "SUBTITLE", "SLIDE_NUMBER"})

    titre = _placeholder_par_type(mise_en_page, "CENTER_TITLE")
    titre.name = "Titre"
    titre.left, titre.top = Emu(CONTENU_L), Emu(2400000)
    titre.width, titre.height = Emu(9000000), Emu(1900000)
    regler_texte(
        titre,
        [dict(taille=36, gras=True, police=POLICE_TITRE, couleur=BLANC,
              interligne=1.06, avant=0, interlettrage=-72)],
        ancrage=MSO_ANCHOR.BOTTOM,
        invite="Titre de couverture : la conclusion, douze mots au plus",
    )
    filet(mise_en_page, CONTENU_L, 4400000, FILET_LARG, LIME)

    sous_titre = _placeholder_par_type(mise_en_page, "SUBTITLE")
    sous_titre.name = "Sous-titre"
    sous_titre.left, sous_titre.top = Emu(CONTENU_L), Emu(4620000)
    sous_titre.width, sous_titre.height = Emu(8000000), Emu(900000)
    regler_texte(
        sous_titre,
        [dict(taille=18, couleur=GRIS_300, interligne=1.5, avant=0)],
        invite="Sous-titre : le cadrage, destinataire et décision à prendre",
    )
    poser_reperes(mise_en_page, GRIS_300)


def construire_section(mise_en_page):
    mise_en_page.name = "Section"
    poser_fond(mise_en_page, NOIR)
    supprimer_placeholders(mise_en_page, {"TITLE", "BODY", "SLIDE_NUMBER"})

    titre = _placeholder_par_type(mise_en_page, "TITLE")
    titre.name = "Titre"
    titre.left, titre.top = Emu(CONTENU_L), Emu(2100000)
    titre.width, titre.height = Emu(9600000), Emu(1400000)
    regler_texte(
        titre,
        [dict(taille=36, gras=True, police=POLICE_TITRE, couleur=BLANC,
              interligne=1.12, avant=0, interlettrage=-72)],
        ancrage=MSO_ANCHOR.BOTTOM,
        invite="Titre de section",
    )
    filet(mise_en_page, CONTENU_L, 3560000, FILET_LARG, LIME)

    corps = _placeholder_par_type(mise_en_page, "BODY")
    corps.name = "Objectif"
    corps.left, corps.top = Emu(CONTENU_L), Emu(3760000)
    corps.width, corps.height = Emu(8600000), Emu(1200000)
    regler_texte(
        corps,
        [dict(taille=18, couleur=GRIS_300, interligne=1.5, avant=0)],
        invite="Objectif de la section : à la fin, vous saurez…",
    )
    poser_reperes(mise_en_page, GRIS_300)


def construire_une_colonne(mise_en_page):
    mise_en_page.name = "Une colonne"
    poser_fond(mise_en_page, PAPIER)
    supprimer_placeholders(mise_en_page, {"TITLE", "OBJECT", "SLIDE_NUMBER"})
    _bandeau_titre(mise_en_page, INVITE_TITRE)

    corps = _placeholder_par_type(mise_en_page, "OBJECT")
    corps.name = "Corps"
    corps.left, corps.top = Emu(CONTENU_L), Emu(CORPS_Y)
    corps.width, corps.height = Emu(CONTENU_LARG), Emu(CORPS_HAUTEUR)
    regler_texte(
        corps, _niveaux_corps(),
        invite="Corps, Inter 18 pt. Quatre puces au plus, deux lignes par puce.",
    )
    poser_reperes(mise_en_page, GRIS_500)


def construire_deux_colonnes(mise_en_page):
    mise_en_page.name = "Deux colonnes"
    poser_fond(mise_en_page, PAPIER)
    supprimer_placeholders(mise_en_page, {"TITLE", "OBJECT", "SLIDE_NUMBER"})
    _bandeau_titre(mise_en_page, INVITE_TITRE)

    for idx, gauche, invite in (
        (1, CONTENU_L, "Colonne de gauche"),
        (2, CONTENU_L + COLONNE_LARG + GOUTTIERE, "Colonne de droite"),
    ):
        colonne = mise_en_page.placeholders.get(idx)
        colonne.name = invite
        colonne.left, colonne.top = Emu(gauche), Emu(CORPS_Y)
        colonne.width, colonne.height = Emu(COLONNE_LARG), Emu(CORPS_HAUTEUR)
        regler_texte(colonne, _niveaux_corps(), invite=invite)
    poser_reperes(mise_en_page, GRIS_500)


def construire_tableau(mise_en_page):
    mise_en_page.name = "Tableau"
    poser_fond(mise_en_page, PAPIER)
    supprimer_placeholders(mise_en_page, {"TITLE", "SLIDE_NUMBER"})
    _bandeau_titre(mise_en_page, INVITE_TITRE)

    # Un vrai espace réservé de tableau : un tableau reste un tableau, jamais
    # une image. Six lignes au plus, en-tête compris.
    ajouter_placeholder(mise_en_page, "Tableau", "tbl", 1,
                        CONTENU_L, CORPS_Y, CONTENU_LARG, CORPS_HAUTEUR)
    poser_reperes(mise_en_page, GRIS_500)


VERBATIM_INTERIEUR = 250000
VERBATIM_SOURCE_Y = CORPS_Y + VERBATIM_INTERIEUR
VERBATIM_FILET_Y = VERBATIM_SOURCE_Y + 300000
VERBATIM_CITATION_Y = VERBATIM_FILET_Y + FILET_EP + 100000


def construire_verbatim(mise_en_page):
    mise_en_page.name = "Verbatim"
    poser_fond(mise_en_page, PAPIER)
    supprimer_placeholders(mise_en_page, {"TITLE", "SLIDE_NUMBER"})
    _bandeau_titre(mise_en_page, INVITE_TITRE)

    cadre_fin(mise_en_page, CONTENU_L, CORPS_Y, CONTENU_LARG, CORPS_HAUTEUR)
    interieur_l = CONTENU_L + VERBATIM_INTERIEUR
    interieur_larg = CONTENU_LARG - 2 * VERBATIM_INTERIEUR

    # L'étiquette de source se lit AVANT la citation, comme dans la charte.
    source = ajouter_placeholder(mise_en_page, "Source du verbatim", "body", 1,
                                 interieur_l, VERBATIM_SOURCE_Y,
                                 interieur_larg, 300000)
    regler_texte(
        source,
        [dict(taille=11, police=POLICE_MONO, couleur=GRIS_700, gras=True,
              interligne=1.2, avant=0, interlettrage=120, capitales="all")],
        invite="Source de la citation",
    )
    filet(mise_en_page, interieur_l, VERBATIM_FILET_Y, 40 * MM, LIME_TRAIT,
          nom="Filet de l'étiquette")

    citation = ajouter_placeholder(
        mise_en_page, "Citation", "body", 2,
        interieur_l, VERBATIM_CITATION_Y, interieur_larg,
        CORPS_Y + CORPS_HAUTEUR - VERBATIM_INTERIEUR - VERBATIM_CITATION_Y,
    )
    regler_texte(
        citation,
        [dict(taille=24, police=POLICE_TITRE, italique=True, couleur=NOIR,
              interligne=1.46, avant=0)],
        invite="Citation, mot pour mot, sans guillemets d'encadrement",
    )
    poser_reperes(mise_en_page, GRIS_500)


# Mise en page par défaut de départ → constructeur.
CONSTRUCTEURS = (
    (0, construire_couverture),     # Title Slide
    (2, construire_section),        # Section Header
    (1, construire_une_colonne),    # Title and Content
    (3, construire_deux_colonnes),  # Two Content
    (5, construire_tableau),        # Title Only
    (4, construire_verbatim),       # Comparison
)


ORDRE = ("Couverture", "Section", "Une colonne", "Deux colonnes",
         "Tableau", "Verbatim")


def _ordonner(masque, noms):
    """PowerPoint présente les mises en page dans l'ordre du `sldLayoutIdLst`,
    pas dans celui des noms : on le remet dans l'ordre du modèle."""
    liste = masque._element.find(qn("p:sldLayoutIdLst"))
    rang = {nom: i for i, nom in enumerate(noms)}
    for entree in sorted(
        list(liste),
        key=lambda e: rang[masque.part.related_part(
            e.get(qn("r:id"))).slide_layout.name],
    ):
        liste.append(entree)


def generer_modele():
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = LARGEUR, HAUTEUR
    ecrire_theme(presentation)

    masque = presentation.slide_masters[0]
    masque.name = "WAOUP"
    preparer_masque(masque)

    mises_en_page = list(masque.slide_layouts)
    gardees = []
    for indice, constructeur in CONSTRUCTEURS:
        constructeur(mises_en_page[indice])
        gardees.append(mises_en_page[indice])
    for mise_en_page in mises_en_page:
        if mise_en_page not in gardees:
            masque.slide_layouts.remove(mise_en_page)  # les inutilisées partent

    _ordonner(masque, ORDRE)

    proprietes = presentation.core_properties
    proprietes.title = "Modèle WAOUP, diapositives 16:9"
    proprietes.author = "WAOUP"
    proprietes.comments = (
        "Généré par scripts/generer_modele.py depuis assets/waoup-charte.css.")

    MODELE.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(MODELE)
    return MODELE


# --------------------------------------------------------------------------- #
# 7. Vérification : on rouvre le fichier et on liste ce qu'il contient
# --------------------------------------------------------------------------- #


def verifier():
    presentation = Presentation(MODELE)
    print(f"Format : {presentation.slide_width / 914400:.3f} × "
          f"{presentation.slide_height / 914400:.3f} in")
    for mise_en_page in presentation.slide_masters[0].slide_layouts:
        print(f"\n{mise_en_page.name}")
        for forme in mise_en_page.placeholders:
            f = forme.placeholder_format
            print(f"    idx {f.idx:>2}  {str(f.type):<18} {forme.name:<22} "
                  f"{forme.left / 914400:5.2f} / {forme.top / 914400:5.2f} in   "
                  f"{forme.width / 914400:5.2f} × {forme.height / 914400:4.2f} in")
        for forme in mise_en_page.shapes:
            if not forme.is_placeholder:
                print(f"    fixe      {forme.name}")


# --------------------------------------------------------------------------- #
# 8. Deck d'exemple : demo/exemple-deck-charte.pptx
#    Contenu repris mot pour mot de demo/01-restitution-avec-skill.md et
#    cas/ardan-transcription-entretien.md. Rien n'est écrit ici.
# --------------------------------------------------------------------------- #


def _poser_numero(diapositive, mise_en_page):
    """python-pptx ne clone pas les espaces réservés latents : le numéro de
    page se pose explicitement sur chaque diapositive."""
    for forme in mise_en_page.placeholders:
        if str(forme.placeholder_format.type).startswith("SLIDE_NUMBER"):
            copie = copy.deepcopy(forme._element)
            for noeud in copie.iter(qn("p:cNvPr")):
                noeud.set("id", "900")
            diapositive.shapes._spTree.append(copie)


def _remplir(porteur, lignes):
    cadre = porteur.text_frame
    cadre.text = lignes[0]
    for ligne in lignes[1:]:
        cadre.add_paragraph().text = ligne


def _en_tete_de_colonne(porteur):
    """Le premier paragraphe d'une colonne est son intitulé : sans puce."""
    paragraphe = porteur.text_frame.paragraphs[0]
    for fragment in paragraphe.runs:
        fragment.font.bold = True
        fragment.font.name = POLICE_TITRE
    p_pr = paragraphe._p.get_or_add_pPr()
    p_pr.set("marL", "0")
    p_pr.set("indent", "0")
    p_pr.append(parse_xml(f'<a:buNone {nsdecls("a")}/>'))


def _styler_cellule(cellule, fond, bordure_bas=None, en_tete=False):
    tc_pr = cellule._tc.get_or_add_tcPr()
    trait = ""
    if bordure_bas:
        trait = (f'<a:lnB w="12700" cap="flat" cmpd="sng" algn="ctr">'
                 f'<a:solidFill><a:srgbClr val="{bordure_bas}"/></a:solidFill>'
                 f'<a:prstDash val="solid"/></a:lnB>')
    tc_pr.getparent().replace(
        tc_pr,
        parse_xml(
            f'<a:tcPr {nsdecls("a")} marL="91440" marR="91440" marT="54864" '
            f'marB="54864" anchor="t">{trait}'
            f'<a:solidFill><a:srgbClr val="{fond}"/></a:solidFill></a:tcPr>'
        ),
    )
    for paragraphe in cellule.text_frame.paragraphs:
        for fragment in paragraphe.runs:
            fragment.font.size = Pt(12 if en_tete else 13)
            fragment.font.bold = en_tete
            fragment.font.name = POLICE_TITRE if en_tete else POLICE_TEXTE
            fragment.font.color.rgb = _rgb(BLANC if en_tete else NOIR)


def _poser_tableau(diapositive, donnees):
    cadre = diapositive.placeholders[1].insert_table(rows=len(donnees), cols=2)
    tableau = cadre.table
    tableau.first_row = False
    tableau.horz_banding = False
    style = tableau._tbl.tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # aucun style
    tableau.columns[0].width = Emu(int(CONTENU_LARG * 0.45))
    tableau.columns[1].width = Emu(CONTENU_LARG - int(CONTENU_LARG * 0.45))
    for rang, ligne in enumerate(donnees):
        for colonne, texte in enumerate(ligne):
            cellule = tableau.cell(rang, colonne)
            cellule.text = texte
            if rang == 0:
                _styler_cellule(cellule, NOIR, en_tete=True)
            else:
                _styler_cellule(cellule, FOND_DOUX if rang % 2 == 0 else PAPIER,
                                bordure_bas=GRIS_200)


def generer_demo():
    presentation = Presentation(MODELE)
    pages = {m.name: m for m in presentation.slide_masters[0].slide_layouts}

    # 1. Couverture — titre et cadrage du document
    diapositive = presentation.slides.add_slide(pages["Couverture"])
    diapositive.shapes.title.text_frame.text = (
        "Restitution, Groupe Ardan, 12 septembre 2026")
    diapositive.placeholders[1].text_frame.text = (
        "Ce que cette restitution doit permettre de décider : le périmètre de "
        "la vague d'entretiens suivante, et l'ordre des chantiers dans la propale.")
    _poser_numero(diapositive, pages["Couverture"])

    # 2. Section — bloc 4
    diapositive = presentation.slides.add_slide(pages["Section"])
    diapositive.shapes.title.text_frame.text = "Synthèse par enjeu"
    diapositive.placeholders[1].text_frame.text = (
        "Cinq enjeux. Chacun en trois colonnes : ce qui se passe, ce que le "
        "terrain propose déjà, la parole exacte.")
    _poser_numero(diapositive, pages["Section"])

    # 3. Une colonne — enjeu 4.4, colonne « Problèmes rencontrés »
    diapositive = presentation.slides.add_slide(pages["Une colonne"])
    diapositive.shapes.title.text_frame.text = (
        "Le client ne pardonne pas l'interlocuteur introuvable")
    _remplir(diapositive.placeholders[1], [
        "La réorganisation a remplacé un contact unique par trois canaux.",
        "Le client ne sait plus vers qui se tourner et le dit à son directeur "
        "d'agence.",
        "L'écart de gravité entre une panne de stock et une panne de relation "
        "est jugé considérable.",
    ])
    _poser_numero(diapositive, pages["Une colonne"])

    # 4. Deux colonnes — enjeu 4.3, problèmes en regard des solutions
    diapositive = presentation.slides.add_slide(pages["Deux colonnes"])
    diapositive.shapes.title.text_frame.text = (
        "La caution immobilise mille à deux mille euros et fait renoncer des "
        "artisans à la machine")
    _remplir(diapositive.placeholders[1], [
        "Problèmes rencontrés",
        "Le montant bloqué est sans rapport avec la durée de location.",
        "Les plus petites structures ne peuvent pas immobiliser cette somme.",
        "Le client renonce à la machine et fait le travail à la main.",
    ])
    _remplir(diapositive.placeholders[2], [
        "Solutions envisagées par le terrain",
        "Aucune piste. L'interviewé attend la nôtre.",
        "Aucune.",
        "Aucune.",
    ])
    for indice in (1, 2):
        _en_tete_de_colonne(diapositive.placeholders[indice])
    _poser_numero(diapositive, pages["Deux colonnes"])

    # 5. Tableau — bloc 6, « Ce qui reste à valider »
    diapositive = presentation.slides.add_slide(pages["Tableau"])
    diapositive.shapes.title.text_frame.text = "Ce qui reste à valider"
    _poser_tableau(diapositive, [
        ("Hypothèse", "Comment la trancher"),
        ("Deux ou trois locations perdues par semaine est le bon ordre de "
         "grandeur",
         "Relevé au comptoir sur quatre semaines, dans plusieurs agences, "
         "avec le motif"),
        ("Le renoncement pour cause de caution est fréquent, au-delà des cas "
         "dont le directeur se souvient",
         "Entretiens artisans, et comptage des devis sans suite"),
        ("Le client capté une fois par le loueur généraliste y reste",
         "Rien dans l'entretien ne le dit. À poser en entretien client"),
        ("Le siège est le décideur du projet",
         "Une seule phrase le suggère, dite en riant. À vérifier avant "
         "d'engager la propale"),
        ("Les clients refusent une application supplémentaire",
         "Dit par un directeur au nom de ses clients. À entendre de la bouche "
         "d'un client"),
    ])
    _poser_numero(diapositive, pages["Tableau"])

    # 6. Verbatim — la phrase source, en regard du titre de la restitution :
    #    « de mes appels » dans l'entretien, « des appels » dans le titre.
    diapositive = presentation.slides.add_slide(pages["Verbatim"])
    diapositive.shapes.title.text_frame.text = (
        "Quatre-vingts pour cent des appels posent la même question : la "
        "machine est-elle là ?")
    diapositive.placeholders[1].text_frame.text = (
        "Directeur d'agence, Groupe Ardan")
    diapositive.placeholders[2].text_frame.text = (
        "Quatre-vingts pour cent de mes appels, c'est ça : « t'as la machine ? »")
    _poser_numero(diapositive, pages["Verbatim"])

    DEMO.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(DEMO)
    return DEMO


# --------------------------------------------------------------------------- #

def main():
    analyseur = argparse.ArgumentParser(description=__doc__)
    analyseur.add_argument("--demo", action="store_true",
                           help="génère aussi demo/exemple-deck-charte.pptx")
    options = analyseur.parse_args()

    print(f"Modèle écrit : {generer_modele()}")
    verifier()
    if options.demo:
        print(f"\nDeck d'exemple : {generer_demo()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
